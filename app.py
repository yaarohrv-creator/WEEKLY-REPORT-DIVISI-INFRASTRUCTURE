import os
import io
import sqlite3
import pandas as pd
import streamlit as st
import re

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    from PIL import Image as PILImage
    from openpyxl.drawing.image import Image as OpenPyXLImage
    has_pil = True
except ImportError:
    st.error("⚠️ Library 'Pillow' belum terinstal.")
    has_pil = False

st.set_page_config(page_title="Sistem Progress Proyek", layout="wide")

if os.path.exists("logo.png"):
    st.sidebar.image("logo.png", use_container_width=True)

UPLOAD_DIR = "uploads"
if not os.path.exists(UPLOAD_DIR):
    os.makedirs(UPLOAD_DIR)

# ---------------------------------------------------------
# AUTENTIKASI PASSWORD
# ---------------------------------------------------------
if "authenticated" not in st.session_state:
    st.session_state["authenticated"] = False

def check_password():
    password_benar = st.secrets.get("APP_PASSWORD", "123456")
    if st.session_state["password_input"] == password_benar:
        st.session_state["authenticated"] = True
        del st.session_state["password_input"]
    else:
        st.session_state["authenticated"] = False
        st.error("🔑 Password salah! Silakan coba lagi.")

if not st.session_state["authenticated"]:
    st.title("🔒 Akses Terbatas - Laporan Progress Proyek")
    st.write("Silakan masukkan password tim untuk mengakses aplikasi.")
    st.text_input("Password Akses:", type="password", key="password_input", on_change=check_password)
    st.info("💡 Silakan hubungi admin untuk mendapatkan password akses.")
    st.stop()

# ---------------------------------------------------------
# DATABASE SQLITE
# ---------------------------------------------------------
def get_db_connection():
    return sqlite3.connect('proyek_v2.db')

def sanitize_filename(filename):
    return re.sub(r'[\\/*?:"<>|]', "", str(filename)).replace(" ", "_")

def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS master_spk (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            no_spk TEXT,
            kontraktor TEXT,
            jenis_pekerjaan TEXT,
            unit TEXT,
            jumlah INTEGER DEFAULT 1,
            nilai_pekerjaan REAL,
            catatan TEXT,
            UNIQUE(no_spk, jenis_pekerjaan)
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS laporan_mingguan (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            waktu_input DATETIME DEFAULT CURRENT_TIMESTAMP,
            no_spk TEXT,
            jenis_pekerjaan TEXT,
            kontraktor TEXT,
            unit TEXT,
            jumlah INTEGER,
            nilai_pekerjaan REAL,
            progress_minggu_lalu REAL,
            progress_minggu_ini REAL,
            catatan TEXT,
            foto_1 TEXT,
            foto_2 TEXT
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS history_progress (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            waktu_input DATETIME DEFAULT CURRENT_TIMESTAMP,
            no_spk TEXT,
            jenis_pekerjaan TEXT,
            kontraktor TEXT,
            unit TEXT,
            progress_minggu_lalu REAL,
            progress_minggu_ini REAL,
            progres_penambahan REAL,
            catatan TEXT,
            foto_1 TEXT,
            foto_2 TEXT
        )
    ''')
    conn.commit()
    conn.close()

init_db()

# ---------------------------------------------------------
# EXPORT EXCEL
# ---------------------------------------------------------
def generate_excel_full_feature(df):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_excel = df.drop(columns=['Hapus', 'real_id'], errors='ignore').copy()
        
        df_progress = df_excel.drop(
            columns=['Foto 1', 'Foto 2', 'Pratinjau Foto 1', 'Pratinjau Foto 2'], 
            errors='ignore'
        ).copy()
        
        try:
            target_col_idx = df_progress.columns.get_loc('Catatan Pekerjaan Terbaru') + 1
            df_progress.insert(target_col_idx, 'Dokumentasi', '') 
        except Exception:
            df_progress['Dokumentasi'] = ''

        df_progress.to_excel(writer, index=False, sheet_name='Laporan Progress')
        
        worksheet_progress = writer.sheets['Laporan Progress']
        workbook = writer.book
        worksheet_foto = workbook.create_sheet(title='Foto Dokumentasi')
        
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border_standard = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        align_center = Alignment(horizontal="center", vertical="center")

        for col_num in range(1, worksheet_progress.max_column + 1):
            cell = worksheet_progress.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_standard

        for row_idx in range(2, worksheet_progress.max_row + 1):
            for col_idx in range(1, worksheet_progress.max_column + 1):
                cell = worksheet_progress.cell(row=row_idx, column=col_idx)
                cell.border = border_standard
                cell.alignment = Alignment(vertical="center")

        for col in worksheet_progress.columns:
            header_name = col[0].value
            if header_name != 'Dokumentasi':
                max_len = max(len(str(cell.value or '')) for cell in col)
                worksheet_progress.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 12)
            else:
                worksheet_progress.column_dimensions[get_column_letter(col[0].column)].width = 15

    return output.getvalue()

# ---------------------------------------------------------
# NAVIGASI
# ---------------------------------------------------------
st.sidebar.title("Navigasi")
menu = st.sidebar.selectbox("Pilih Menu", [
    "Dashboard Progress", 
    "Input Progress Mingguan",
    "Kelola Master SPK"
])

st.sidebar.markdown("---")
if st.sidebar.button("🚪 Logout"):
    st.session_state["authenticated"] = False
    st.rerun()

# ---------------------------------------------------------
# MENU 1: DASHBOARD PROGRESS
# ---------------------------------------------------------
if menu == "Dashboard Progress":
    st.title("📊 WEEKLY REPORT DIVISI INFRASTRUCTURE")

    tab_bangka, tab_belitung, tab_semua, tab_history = st.tabs([
        "🏝️ Laporan Progress Bangka", 
        "🏖️ Laporan Progress Belitung", 
        "📋 Semua Progress Proyek", 
        "📜 Riwayat / History Update"
    ])

    query_view = """
        SELECT 
            id AS real_id,
            waktu_input AS [Waktu Input Terbaru],
            no_spk AS [Nomor SPK],
            kontraktor AS [Nama Kontraktor],
            jenis_pekerjaan AS [Jenis Pekerjaan],
            unit AS [Unit Proyek],
            CAST(COALESCE(jumlah, 1) AS INTEGER) AS [Jumlah],
            nilai_pekerjaan AS [Nilai Kontrak Pekerjaan Ini (Rp)],
            progress_minggu_lalu AS [Progress Minggu Lalu (%)],
            progress_minggu_ini AS [Progress Minggu Ini (%)],
            (COALESCE(progress_minggu_ini, 0) - COALESCE(progress_minggu_lalu, 0)) AS [Selisih / Varian (%)],
            catatan AS [Catatan Pekerjaan Terbaru],
            foto_1 AS [Pratinjau Foto 1],
            foto_2 AS [Pratinjau Foto 2]
        FROM laporan_mingguan
        ORDER BY id ASC
    """

    with get_db_connection() as conn:
        try:
            df_all = pd.read_sql_query(query_view, conn)
        except Exception:
            df_all = pd.DataFrame()

    def render_dashboard_table(df_data, tab_key_prefix):
        if df_data.empty:
            st.info("💡 Belum ada data progress untuk wilayah/kategori ini.")
            return

        # Reset index agar index df_data dan edited_df sinkron sempurna (0, 1, 2, ...)
        df_data = df_data.reset_index(drop=True)

        df_display = df_data.copy()
        df_display.insert(0, "Hapus", False)
        if 'No' not in df_display.columns:
            df_display.insert(1, 'No', range(1, len(df_display) + 1))

        column_order = [
            'Hapus', 'No', 'Nomor SPK', 'Nama Kontraktor', 'Jenis Pekerjaan', 'Unit Proyek', 'Jumlah',
            'Nilai Kontrak Pekerjaan Ini (Rp)', 'Progress Minggu Lalu (%)', 'Progress Minggu Ini (%)',
            'Selisih / Varian (%)', 'Catatan Pekerjaan Terbaru', 'Pratinjau Foto 1', 'Pratinjau Foto 2'
        ]
        existing_cols = [c for c in column_order if c in df_display.columns]

        # Tampilan Tabel Interaktif
        edited_df = st.data_editor(
            df_display[existing_cols],
            use_container_width=True,
            hide_index=True,
            column_config={
                "Hapus": st.column_config.CheckboxColumn("Hapus?", help="Centang untuk menghapus baris ini"),
                "Jumlah": st.column_config.NumberColumn("Jumlah", format="%d"),
                "Nilai Kontrak Pekerjaan Ini (Rp)": st.column_config.NumberColumn("Nilai Kontrak (Rp)", format="Rp %d"),
                "Progress Minggu Lalu (%)": st.column_config.NumberColumn("Progress Minggu Lalu (%)", format="%.2f %%"),
                "Progress Minggu Ini (%)": st.column_config.NumberColumn("Progress Minggu Ini (%)", format="%.2f %%"),
                "Selisih / Varian (%)": st.column_config.NumberColumn("Selisih / Varian (%)", format="%.2f %%"),
                "Pratinjau Foto 1": st.column_config.ImageColumn("Pratinjau Foto 1"),
                "Pratinjau Foto 2": st.column_config.ImageColumn("Pratinjau Foto 2"),
            },
            key=f"editor_{tab_key_prefix}"
        )

        col_sav, col_del_btn = st.columns([2, 2])
        
        # TOMBOL SIMPAN
        with col_sav:
            if st.button("💾 Simpan Edit Progress", key=f"btn_save_{tab_key_prefix}"):
                with get_db_connection() as conn:
                    cursor = conn.cursor()
                    for idx, row in edited_df.iterrows():
                        real_id = df_data.loc[idx, 'real_id']
                        cursor.execute("""
                            UPDATE laporan_mingguan
                            SET progress_minggu_ini = ?, catatan = ?
                            WHERE id = ?
                        """, (row.get('Progress Minggu Ini (%)'), row.get('Catatan Pekerjaan Terbaru'), real_id))
                    conn.commit()
                st.success("Perubahan data berhasil disimpan!")
                st.rerun()

        # TOMBOL HAPUS BARIS CENTANG (PASTI TERHAPUS BERSIH & BEBAS ERROR)
        with col_del_btn:
            if st.button("🗑️ Hapus Baris Yang Dicentang", key=f"btn_del_chk_{tab_key_prefix}", type="primary"):
                rows_to_delete = edited_df[edited_df["Hapus"] == True]
                if rows_to_delete.empty:
                    st.warning("⚠️ Centang kolom 'Hapus?' pada baris yang ingin Anda hapus terlebih dahulu!")
                else:
                    with get_db_connection() as conn:
                        cursor = conn.cursor()
                        for idx in rows_to_delete.index:
                            target_id = df_data.loc[idx, 'real_id']
                            spk_num = df_data.loc[idx, 'Nomor SPK']
                            j_pekerjaan = df_data.loc[idx, 'Jenis Pekerjaan']
                            
                            # Hapus dari laporan_mingguan
                            cursor.execute("DELETE FROM laporan_mingguan WHERE id = ?", (target_id,))
                            # Hapus dari master_spk
                            cursor.execute("DELETE FROM master_spk WHERE no_spk = ? AND jenis_pekerjaan = ?", (spk_num, j_pekerjaan))
                        
                        conn.commit()
                    st.success("✅ Baris terpilih berhasil dihapus permanen dari Database!")
                    st.rerun()

        st.markdown("---")
        st.subheader("📥 Export & Download Laporan")
        try:
            excel_bytes = generate_excel_full_feature(df_display)
            st.download_button(
                label=f"📥 Download Laporan ({tab_key_prefix.capitalize()}) - Excel",
                data=excel_bytes,
                file_name=f"Laporan_Progress_{tab_key_prefix}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key=f"dl_{tab_key_prefix}"
            )
        except Exception as e:
            st.error(f"Gagal memproses file Excel: {e}")

    # TAB BANGKA
    with tab_bangka:
        st.subheader("📍 Laporan Progress Proyek - Wilayah Bangka")
        if not df_all.empty:
            df_bangka = df_all[df_all['Unit Proyek'].astype(str).str.contains('BANGKA|BKA', case=False, na=False)]
            render_dashboard_table(df_bangka, "bangka")

    # TAB BELITUNG
    with tab_belitung:
        st.subheader("📍 Laporan Progress Proyek - Wilayah Belitung")
        if not df_all.empty:
            pola_belitung = 'BELITUNG|BLT|BPSL|BPRE|BPT'
            df_belitung = df_all[
                df_all['Unit Proyek'].astype(str).str.contains(pola_belitung, case=False, na=False) |
                (~df_all['Unit Proyek'].astype(str).str.contains('BANGKA|BKA', case=False, na=False))
            ]
            render_dashboard_table(df_belitung, "belitung")

    # TAB SEMUA
    with tab_semua:
        st.subheader("🌐 Semua Laporan Progress Proyek")
        render_dashboard_table(df_all, "semua")

    # TAB RIWAYAT
    with tab_history:
        st.subheader("📜 Log Riwayat Update")
        with get_db_connection() as conn:
            df_history = pd.read_sql_query("SELECT * FROM history_progress ORDER BY waktu_input DESC", conn)
        
        st.dataframe(
            df_history, 
            use_container_width=True,
            column_config={
                "progress_minggu_lalu": st.column_config.NumberColumn(format="%.2f %%"),
                "progress_minggu_ini": st.column_config.NumberColumn(format="%.2f %%"),
                "progres_penambahan": st.column_config.NumberColumn(format="%.2f %%"),
            }
        )

# ---------------------------------------------------------
# MENU 2: INPUT PROGRESS
# ---------------------------------------------------------
elif menu == "Input Progress Mingguan":
    st.title("📝 Input Laporan Progress Mingguan Berdasarkan SPK")
    st.markdown("---")

    with get_db_connection() as conn:
        df_master_all = pd.read_sql_query("SELECT * FROM master_spk", conn)

    if df_master_all.empty:
        st.warning("⚠️ Belum ada data Master SPK.")
    else:
        tab_i_bangka, tab_i_belitung = st.tabs(["🏝️ Input Bangka", "🏖️ Input Belitung"])

        def render_input_form(df_master_wilayah, tab_key_prefix):
            if df_master_wilayah.empty:
                st.info("💡 Belum ada data master untuk wilayah ini.")
                return

            list_spk_unique = sorted(df_master_wilayah['no_spk'].unique().tolist())
            selected_spk_no = st.selectbox("Pilih Nomor SPK:", list_spk_unique, key=f"select_spk_no_{tab_key_prefix}")

            df_spk_filtered = df_master_wilayah[df_master_wilayah['no_spk'] == selected_spk_no]
            list_pekerjaan = df_spk_filtered['jenis_pekerjaan'].unique().tolist()
            selected_pekerjaan = st.selectbox("Pilih Jenis Pekerjaan:", list_pekerjaan, key=f"select_pekerjaan_{tab_key_prefix}")

            spk_data_selected = df_spk_filtered[df_spk_filtered['jenis_pekerjaan'] == selected_pekerjaan].iloc[0]

            with st.form(f"form_input_week_{tab_key_prefix}", clear_on_submit=True):
                prog_ini = st.number_input("Progress Akumulatif Minggu Ini (%)", min_value=0.0, max_value=100.0, value=0.0, step=0.01)
                catatan_lap = st.text_area("Catatan Pekerjaan Minggu Ini")
                f_upload_1 = st.file_uploader("Upload Foto 1", type=["jpg", "jpeg", "png"], key=f"f1_{tab_key_prefix}")
                f_upload_2 = st.file_uploader("Upload Foto 2", type=["jpg", "jpeg", "png"], key=f"f2_{tab_key_prefix}")

                if st.form_submit_button("💾 Simpan Laporan"):
                    path_f1, path_f2 = None, None
                    import time
                    ts = int(time.time())
                    if f_upload_1:
                        path_f1 = os.path.join(UPLOAD_DIR, f"{ts}_f1.jpg")
                        with open(path_f1, "wb") as f: f.write(f_upload_1.getbuffer())
                    if f_upload_2:
                        path_f2 = os.path.join(UPLOAD_DIR, f"{ts}_f2.jpg")
                        with open(path_f2, "wb") as f: f.write(f_upload_2.getbuffer())

                    with get_db_connection() as conn:
                        cursor = conn.cursor()
                        cursor.execute("DELETE FROM laporan_mingguan WHERE no_spk=? AND jenis_pekerjaan=?", (spk_data_selected['no_spk'], spk_data_selected['jenis_pekerjaan']))
                        cursor.execute("""
                            INSERT INTO laporan_mingguan (
                                no_spk, jenis_pekerjaan, kontraktor, unit, jumlah, nilai_pekerjaan, 
                                progress_minggu_lalu, progress_minggu_ini, catatan, foto_1, foto_2, waktu_input
                            ) VALUES (?,?,?,?,?,?,?,?,?,?,?, CURRENT_TIMESTAMP)""", 
                            (spk_data_selected['no_spk'], spk_data_selected['jenis_pekerjaan'], spk_data_selected['kontraktor'], spk_data_selected['unit'], int(spk_data_selected['jumlah'] or 1), spk_data_selected['nilai_pekerjaan'], 0.0, prog_ini, catatan_lap, path_f1, path_f2))
                        conn.commit()
                    st.success("✅ Laporan berhasil disimpan!")
                    st.rerun()

        with tab_i_bangka:
            df_bangka_master = df_master_all[df_master_all['unit'].astype(str).str.contains('BANGKA|BKA', case=False, na=False)]
            render_input_form(df_bangka_master, "bangka")

        with tab_i_belitung:
            pola_belitung = 'BELITUNG|BLT|BPSL|BPRE|BPT'
            df_belitung_master = df_master_all[
                df_master_all['unit'].astype(str).str.contains(pola_belitung, case=False, na=False) |
                (~df_master_all['unit'].astype(str).str.contains('BANGKA|BKA', case=False, na=False))
            ]
            render_input_form(df_belitung_master, "belitung")

# ---------------------------------------------------------
# MENU 3: KELOLA MASTER
# ---------------------------------------------------------
elif menu == "Kelola Master SPK":
    st.title("⚙️ Kelola Master Data Pekerjaan / SPK")
    with st.form("form_tambah_master", clear_on_submit=True):
        wilayah = st.selectbox("Wilayah / Unit Proyek:", ["BANGKA", "BELITUNG"])
        unit_proyek = st.text_input("Nama Unit / Detail Lokasi Proyek:", value=f"UNIT {wilayah}")
        no_spk = st.text_input("Nomor SPK:")
        kontraktor = st.text_input("Nama Kontraktor:")
        jenis_pekerjaan = st.text_input("Jenis Pekerjaan:")
        jumlah = st.number_input("Jumlah Unit/Pekerjaan:", min_value=1, value=1, step=1)
        nilai_pekerjaan = st.number_input("Nilai Kontrak Pekerjaan (Rp):", min_value=0.0, value=0.0, step=1000000.0)

        if st.form_submit_button("➕ Tambahkan ke Master Data"):
            if no_spk and jenis_pekerjaan and kontraktor:
                with get_db_connection() as conn:
                    cursor = conn.cursor()
                    cursor.execute("""
                        INSERT INTO master_spk (no_spk, kontraktor, jenis_pekerjaan, unit, jumlah, nilai_pekerjaan)
                        VALUES (?, ?, ?, ?, ?, ?)
                    """, (no_spk.strip(), kontraktor.strip(), jenis_pekerjaan.strip(), unit_proyek.strip(), int(jumlah), nilai_pekerjaan))
                    conn.commit()
                st.success("✅ Master SPK berhasil ditambahkan!")
                st.rerun()
